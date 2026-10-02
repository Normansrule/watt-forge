"""Generate the flagship loss-model spreadsheet (live formulas), the Jupyter notebook,
and the ngspice cross-check netlist.

    python scripts/make_artifacts.py
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import nbformat  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from watt_forge import devices, magnetics, spice  # noqa: E402
from watt_forge.flagship import model, sim  # noqa: E402
from watt_forge.flagship.params import Params  # noqa: E402

OUT = ROOT / "hardware" / "flagship"


def spreadsheet(path: Path, vin=56.0, vout=48.0, pout=400.0, fs=75e3):
    p = Params()
    dev = devices.get(p.switch)
    bds = devices.get(p.bds)
    mat = magnetics.MATERIALS[p.core]
    wb = Workbook()
    ws = wb.active
    ws.title = "Buck-mode loss model"
    bold = Font(bold=True)
    inp = PatternFill("solid", fgColor="E5EFFF")
    ws["A1"] = "Watt Forge flagship: three-level buck-mode loss budget (live formulas). REFERENCE ONLY."
    ws["A1"].font = bold
    ws["A2"] = "Blue cells are inputs (change them). Closed-form versions of watt_forge/flagship/model.py for D > 0.5 buck mode."
    rows = [
        ("Vin", vin, "V"), ("Vout", vout, "V"), ("Pout", pout, "W"), ("fs", fs, "Hz"), ("L", p.l, "H"),
        ("Rds_max_25C", dev.rds_max_mohm * 1e-3, "ohm"), ("Rds_tc", dev.rds_tc, "1/K"), ("Tj", p.tj, "C"), ("dyn_rds", p.dyn_rds, "x"),
        ("DCR_25C", p.dcr, "ohm"), ("DCR_tc", p.dcr_tc, "1/K"), ("T_ind", p.t_ind, "C"), ("Rac_factor", p.rac_factor, "x"),
        ("Qoss_ref", dev.qoss_nc * 1e-9, "C"), ("Vref_qoss", dev.qoss_v, "V"), ("m_qoss", dev.qoss_exp, "-"),
        ("Qg", dev.qg(), "C"), ("Vdrv", dev.vdrv, "V"), ("Vsd", dev.vsd, "V"), ("t_dead", p.t_dead, "s"),
        ("t_ov_on", dev.t_overlap(True), "s"), ("t_ov_off", dev.t_overlap(False), "s"),
        ("k_hyst", p.coss_hyst_frac, "-"), ("L_loop", p.l_loop, "H"),
        ("ESR_fly", p.cfly_esr(), "ohm"), ("ESR_in", p.esr_in, "ohm"), ("ESR_out", p.esr_out, "ohm"),
        ("R_pcb", p.r_pcb, "ohm"), ("R_shunt", p.r_shunt, "ohm"),
        ("R_bds_25C", bds.rds_max_mohm * 1e-3, "ohm"), ("P_aux", p.p_aux_switching, "W"),
        ("k_B", p.k_b, "T/A"), ("V_core", p.core_volume, "m^3"),
        ("st_alpha", mat["alpha"], "-"), ("st_beta", mat["beta"], "-"),
        ("st_k", magnetics.steinmetz_k(mat), "W/m^3/Hz^a/T^b"),
        ("igse_ki", magnetics.igse_ki(magnetics.steinmetz_k(mat), mat["alpha"], mat["beta"]), "iGSE k_i from st_k, alpha, beta (magnetics.igse_ki)"),
    ]
    ws["A4"], ws["B4"], ws["C4"] = "Input", "Value", "Unit"
    for c in ("A4", "B4", "C4"):
        ws[c].font = bold
    ref = {}
    r = 5
    for name, val, unit in rows:
        ws.cell(r, 1, name)
        ws.cell(r, 2, val).fill = inp
        ws.cell(r, 3, unit)
        ref[name] = f"$B${r}"
        r += 1
    R = lambda n: ref[n]  # noqa: E731
    r += 1
    ws.cell(r, 1, "Derived").font = bold
    r += 1
    derived = [
        ("R_sw", f"={R('Rds_max_25C')}*(1+{R('Rds_tc')}*({R('Tj')}-25))*{R('dyn_rds')}", "ohm"),
        ("DCR_hot", f"={R('DCR_25C')}*(1+{R('DCR_tc')}*({R('T_ind')}-25))", "ohm"),
        ("R_bds", f"={R('R_bds_25C')}*(1+{R('Rds_tc')}*({R('Tj')}-25))", "ohm"),
        ("Iout", f"={R('Pout')}/{R('Vout')}", "A"),
        ("IL", "=Iout_", "A"),
        ("D", f"=({R('Vout')}+IL_*(4*R_sw_+DCR_hot_+{R('R_pcb')}))/{R('Vin')}", "-"),
        ("dI", f"=({R('Vin')}-{R('Vout')})*(D_-0.5)/({R('fs')}*{R('L')})", "A pk-pk"),
        ("I2", "=IL_^2+dI_^2/12", "A^2"),
        ("Iv", "=IL_-dI_/2", "A"),
        ("Ipk", "=IL_+dI_/2", "A"),
        ("Vsw", f"={R('Vin')}/2", "V"),
        ("Qoss_V", f"={R('Qoss_ref')}*(Vsw_/{R('Vref_qoss')})^{R('m_qoss')}", "C"),
        ("Eoss_V", f"={R('m_qoss')}/({R('m_qoss')}+1)*Qoss_V_*Vsw_", "J"),
        ("zvs_res", f"=MAX(0,1-Ipk_*{R('t_dead')}/(2*Qoss_V_))", "-"),
        ("Iin", "=D_*IL_", "A"),
        ("d_core", "=2*D_-1", "-"),
    ]
    names = {}
    for name, formula, unit in derived:
        names[name] = f"$B${r}"
        ws.cell(r, 1, name)
        ws.cell(r, 3, unit)
        r += 1
    # resolve placeholder names (Name_) to cell refs
    rr = r - len(derived)
    for i, (name, formula, unit) in enumerate(derived):
        f = formula
        for n2, cell in names.items():
            f = f.replace(n2 + "_", cell)
        ws.cell(rr + i, 2, f)
    N = lambda n: names[n]  # noqa: E731
    r += 1
    ws.cell(r, 1, "Loss term").font = bold
    ws.cell(r, 2, "Watts").font = bold
    ws.cell(r, 3, "Python model").font = bold
    r += 1
    py = model.evaluate(vin, vout, pout, fs, "buck", p)["losses"]
    terms = [
        ("switch_conduction", f"=4*{N('R_sw')}*{N('I2')}"),
        ("inductor_dcr", f"={N('DCR_hot')}*{N('I2')}"),
        ("inductor_ac", f"={N('DCR_hot')}*({R('Rac_factor')}-1)*{N('dI')}^2/12"),
        ("pcb_copper", f"={R('R_pcb')}*{N('I2')}"),
        ("shunts", f"={R('R_shunt')}*({N('Iin')}^2+{N('Iout')}^2)"),
        ("bds_disconnect", f"={N('R_bds')}*{N('Iin')}^2"),
        ("cfly_esr", f"={R('ESR_fly')}*2*(1-{N('D')})*{N('I2')}"),
        ("cin_cout_esr", f"={R('ESR_in')}*({N('D')}*{N('I2')}-{N('Iin')}^2)+{R('ESR_out')}*{N('dI')}^2/12"),
        ("coss", f"=(2*{N('Qoss_V')}*{N('Vsw')}+2*{N('Qoss_V')}*{N('Vsw')}*{N('zvs_res')}^2)*{R('fs')}"),
        ("overlap", f"=(2*0.5*{N('Vsw')}*{N('Iv')}*{R('t_ov_on')}+2*0.5*{N('Vsw')}*{N('Ipk')}*{R('t_ov_off')})*{R('fs')}"),
        ("dead_time", f"=(2*{R('Vsd')}*{N('Iv')}*{R('t_dead')}+2*{R('Vsd')}*{N('Ipk')}*MAX(0,{R('t_dead')}-2*{N('Qoss_V')}/{N('Ipk')}))*{R('fs')}"),
        ("coss_hysteresis", f"=4*{R('k_hyst')}*{N('Eoss_V')}*{R('fs')}"),
        ("loop_ringing", f"=2*0.5*{R('L_loop')}*{N('Iv')}^2*{R('fs')}"),
        ("gate", f"=4*{R('Qg')}*{R('Vdrv')}*{R('fs')}"),
        ("inductor_core", f"={R('igse_ki')}*({R('k_B')}*{N('dI')})^{R('st_beta')}*(2*{R('fs')})^{R('st_alpha')}*({N('d_core')}^(1-{R('st_alpha')})+(1-{N('d_core')})^(1-{R('st_alpha')}))*{R('V_core')}"),
        ("aux", f"={R('P_aux')}"),
    ]
    first = r
    for name, f in terms:
        ws.cell(r, 1, name)
        ws.cell(r, 2, f)
        ws.cell(r, 3, round(py.get(name, 0.0), 6))
        r += 1
    ws.cell(r, 1, "TOTAL").font = bold
    ws.cell(r, 2, f"=SUM(B{first}:B{r-1})")
    ws.cell(r, 3, f"=SUM(C{first}:C{r-1})")
    r += 1
    ws.cell(r, 1, "Efficiency").font = bold
    ws.cell(r, 2, f"={R('Pout')}/({R('Pout')}+B{r-1})")
    ws.cell(r, 3, f"={R('Pout')}/({R('Pout')}+C{r-1})")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    wb.save(path)
    return {"first": first, "terms": [t for t, _ in terms], "total_row": r - 1}


def notebook(path: Path):
    nb = nbformat.v4.new_notebook()
    cells = [
        ("md", "# Flagship loss model notebook\n\nREFERENCE ONLY: simulate and review before use.\n\nThis notebook drives the same code the tests check: the analytic model, the time-domain circuit simulation, and the efficiency curve. Run it from the repo root after `pip install -e .[dev]`."),
        ("code", "from watt_forge.flagship import model, sim, Params\nfrom watt_forge.style import group_losses\np = Params()\nr = model.best(56.0, 48.0, 400.0, p)\nprint(r['label'], r['fs'], round(100*r['eta'], 3))\nfor k, v in sorted(r['losses'].items(), key=lambda kv: -kv[1]):\n    print(f'{k:20s} {v:8.3f} W')"),
        ("md", "## Circuit simulation vs analytic model\nThe simulator solves the periodic steady state of the real switched circuit (dead time, reverse conduction, unbalanced flying capacitors) and must agree with the model to within 0.15 points."),
        ("code", "for vin, frac in [(12, 1), (24, 1), (36, .5), (48, 1), (56, 1), (60, .1)]:\n    pout = frac * p.spec.p_rated(vin)\n    a = model.best(vin, 48.0, pout, p)\n    s = sim.simulate(vin, 48.0, pout, a['fs'], a['mode'], p, a['phase'])\n    print(f\"{vin:5.1f} V {a['label']:24s} model {100*a['eta']:.3f} %  sim {100*s['eta']:.3f} %\")"),
        ("md", "## Efficiency curve vs input voltage"),
        ("code", "import matplotlib.pyplot as plt\nvins = list(range(12, 61, 2))\nfor frac in (0.1, 0.3, 0.5, 1.0):\n    plt.plot(vins, [100*model.best(v, 48.0, frac*p.spec.p_rated(v), p)['eta'] for v in vins], label=f'{int(frac*100)} %')\nplt.xlabel('Vin [V]'); plt.ylabel('efficiency [%]'); plt.legend(); plt.grid(True); plt.show()"),
        ("md", "## Change a parameter and watch the budget move\nTry `p.switch = 'EPC2302'`, `p.t_dead = 20e-9` or `p.l = 2.2e-6` and re-run."),
        ("code", "from dataclasses import replace\nfor dev in ['EPC2361', 'EPC2302', 'EPC2088', 'ISC030N10NM6']:\n    q = replace(p, switch=dev)\n    r = model.best(36.0, 48.0, 300.0, q)\n    print(dev, round(100*r['eta'], 3), {k: round(v, 3) for k, v in group_losses(r['losses'])})"),
    ]
    for kind, src in cells:
        nb.cells.append(nbformat.v4.new_markdown_cell(src) if kind == "md" else nbformat.v4.new_code_cell(src))
    nbformat.write(nb, str(path))


def netlist(path: Path):
    p = Params()
    a = model.best(56.0, 48.0, 400.0, p)
    s = sim.ideal_switch_efficiency(56.0, 48.0, 400.0, a["fs"], a["mode"], p, a["phase"])
    ckt = s["res"]["ckt"]
    net = spice.flagship_netlist(56.0, 48.0, a["mode"], s["d1"], s["d2"], s["phase"], a["fs"], ckt.r_sw,
                                 ckt.dcr + p.r_pcb, p.l, ckt.c1, ckt.c2, ckt.esr, s["x0"])
    path.write_text(net)
    return s["eta"]


if __name__ == "__main__":
    (ROOT / "notebooks").mkdir(exist_ok=True)
    (OUT / "spice").mkdir(exist_ok=True)
    spreadsheet(OUT / "loss_model.xlsx")
    notebook(ROOT / "notebooks" / "loss_model.ipynb")
    eta = netlist(OUT / "spice" / "flagship_buck_56V_400W.cir")
    print("artifacts written; conduction-only eta for the netlist:", round(eta, 6))
