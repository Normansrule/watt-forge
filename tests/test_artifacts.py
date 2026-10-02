"""Design deliverables exist and agree with the model: spreadsheet, notebook, netlist, BOM, references."""
import csv
import json
from pathlib import Path

import pytest

from watt_forge.flagship import model
from watt_forge.flagship.params import Params

ROOT = Path(__file__).resolve().parents[1]
HW = ROOT / "hardware" / "flagship"


def test_spreadsheet_formulas_match_model():
    pycel = pytest.importorskip("pycel")
    import openpyxl
    xl = pycel.ExcelCompiler(filename=str(HW / "loss_model.xlsx"))
    ws = openpyxl.load_workbook(HW / "loss_model.xlsx").active
    sheet = ws.title
    rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
    total = xl.evaluate(f"{sheet}!B{rows['TOTAL']}")
    py = model.evaluate(56.0, 48.0, 400.0, 75e3, "buck", Params())
    assert total == pytest.approx(py["loss_total"], rel=0.01)
    for term in ("switch_conduction", "coss", "gate", "bds_disconnect"):
        assert xl.evaluate(f"{sheet}!B{rows[term]}") == pytest.approx(py["losses"][term], rel=0.01)


def test_notebook_and_netlist_present():
    nb = json.loads((ROOT / "notebooks" / "loss_model.ipynb").read_text())
    assert nb["nbformat"] == 4 and len(nb["cells"]) >= 6
    net = (HW / "spice" / "flagship_buck_56V_400W.cir").read_text()
    assert ".tran" in net and ".control" not in net


def test_bom_rows_are_traceable():
    rows = list(csv.DictReader(open(HW / "BOM.csv", newline="")))
    assert any(r["part_number"] == "EPC2361" for r in rows)
    for r in rows:
        assert r["status"], r
        if r["status"] != "SELECT":
            assert r["manufacturer"] and r["source"].startswith("https://"), r


def test_references_carry_conditions():
    refs = json.loads((ROOT / "data" / "references.json").read_text())
    ids = set()
    for r in refs["references"]:
        assert r["id"] not in ids
        ids.add(r["id"])
        assert r.get("doi") or r.get("url")
        assert r["status"]
        if r.get("headline") and "%" in r["headline"]:
            assert r.get("conditions"), f"{r['id']}: every efficiency figure needs its test conditions"
    md = (ROOT / "docs" / "REFERENCES.md").read_text()
    assert refs["last_verified"] in md


def test_reference_code_is_labelled():
    for f in ("control/wf_ctrl.c", "control/wf_ctrl.h", "hdl/wf_pwm3l.v"):
        assert "REFERENCE ONLY" in (HW / f).read_text()
